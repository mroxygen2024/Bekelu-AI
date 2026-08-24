from docx import Document
from pathlib import Path
from app.core.logging import logger

class DocxParser:
    def extract_text(self, file_path: Path) -> str:
        try:
            doc = Document(file_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    text += f"\n{row_text}"
            
            logger.info(f"Extracted {len(text)} characters from DOCX")
            return text.strip()
        except Exception as e:
            logger.error(f"DOCX parsing error: {e}")
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

docx_parser = DocxParser()
